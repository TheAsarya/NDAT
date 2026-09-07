"""Build the weekly linebacker-role Excel workbook."""

from __future__ import annotations

import argparse
import json
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

import polars as pl
from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from ndat.config import DataConfig
from ndat.lb_roles import SOURCE_DATASETS, build_season, normalize_threshold
from ndat.manager import DataManager


DEFAULT_PRIMARY_THRESHOLD = 0.85
DEFAULT_COMPARISON_THRESHOLD = 0.80
DEFAULT_TOP_RANK = 16
DEFAULT_SECOND_RANK = 32
DEFAULT_RECENT_WEEKS = 3
DEFAULT_REGULAR_SEASON_END = 18

NAVY = "172C45"
INK = "17202A"
MUTED = "637083"
LINE = "DCE2E8"
TEAM_FILL = "E9EFF7"
WEEKLY_FILLS = {"top": "DFF4E7", "second": "FFF0C2", "below": "E8F2FF"}
SEASON_FILLS = {"top": "674EA7", "second": "D9D2E9"}
RECENT_FILLS = {"top": "0F766E", "second": "CCFBF1"}


def _key(row: Mapping[str, Any]) -> tuple[str, str]:
    return str(row["team"]), str(row["player_id"])


def _competition_ranks(
    players: Sequence[Mapping[str, Any]], field: str
) -> dict[tuple[str, str], int]:
    scores = sorted((float(player[field]) for player in players), reverse=True)
    return {_key(player): scores.index(float(player[field])) + 1 for player in players}


def _rank_band(rank: int, top_rank: int, second_rank: int) -> str:
    if rank <= top_rank:
        return "top"
    if rank <= second_rank:
        return "second"
    return "below"


def _matrix_payload(
    frame: pl.DataFrame, *, recent_week_count: int, regular_season_end: int
) -> dict[str, Any]:
    if frame.is_empty():
        raise ValueError("Cannot build a workbook from an empty linebacker-role frame")
    regular = frame.filter(pl.col("week") <= regular_season_end)
    recent_weeks = sorted(regular["week"].unique().to_list())[-recent_week_count:]
    qualifying = frame.filter(pl.col("qualifies_full_time"))
    players = (
        qualifying.select("team", "player_id", "player_name")
        .unique()
        .sort(["team", "player_name"])
        .to_dicts()
    )
    prepared_players: list[dict[str, Any]] = []
    for player in players:
        games = regular.filter(
            (pl.col("team") == player["team"])
            & (pl.col("player_id") == player["player_id"])
        )
        prepared_players.append(
            player
            | {
                "season_points": games["fantasy_points"].sum(),
                "recent_points": games.filter(pl.col("week").is_in(recent_weeks))[
                    "fantasy_points"
                ].sum(),
            }
        )
    cells = {
        (str(row["team"]), str(row["player_id"]), int(row["week"])): row
        for row in qualifying.to_dicts()
    }
    return {
        "threshold": float(frame["role_threshold"][0]),
        "weeks": sorted(frame["week"].unique().to_list()),
        "recent_weeks": recent_weeks,
        "players": prepared_players,
        "cells": cells,
    }


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=f"FF{color}")


def _apply_font(cell: Cell, *, bold: bool = False, color: str = INK, size: int = 10) -> None:
    cell.font = Font(name="Arial", size=size, bold=bold, color=f"FF{color}")


def _add_matrix_sheet(
    workbook: Workbook,
    payload: Mapping[str, Any],
    *,
    season: int,
    top_rank: int,
    second_rank: int,
) -> None:
    threshold = float(payload["threshold"])
    worksheet = workbook.create_sheet(f"{threshold:.0%} Roles")
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "C4"

    weeks = list(payload["weeks"])
    players = list(payload["players"])
    cells = payload["cells"]
    headers = ["Team", "Player", *(f"W{week}" for week in weeks)]
    final_column = get_column_letter(len(headers))
    worksheet["A1"] = f"{season} linebacker snap-share roles at {threshold:.0%}"
    recent_label = "–".join(str(week) for week in payload["recent_weeks"])
    worksheet["A2"] = (
        f"Weekly cells: green top {top_rank}, amber {top_rank + 1}–{second_rank}, "
        f"blue below {second_rank}. Player: purple regular-season rank. Team: teal "
        f"NFL weeks {recent_label} rank. Dark = top {top_rank}; light = "
        f"{top_rank + 1}–{second_rank}; ties share rank. Missing games = 0; "
        "playoffs excluded from season/recent highlights. Stuff = TFL proxy. "
        "Source: NFLverse via NDAT."
    )
    for column, header in enumerate(headers, start=1):
        worksheet.cell(3, column, header)

    weekly_ranks: dict[tuple[str, str, int], int] = {}
    for week in weeks:
        week_cells = [(key, row) for key, row in cells.items() if key[2] == week]
        scores = sorted(
            (float(row["fantasy_points"]) for _, row in week_cells), reverse=True
        )
        for key, row in week_cells:
            weekly_ranks[key] = scores.index(float(row["fantasy_points"])) + 1

    season_ranks = _competition_ranks(players, "season_points")
    recent_ranks = _competition_ranks(players, "recent_points")
    previous_team: str | None = None
    for row_number, player in enumerate(players, start=4):
        team = str(player["team"])
        worksheet.cell(row_number, 1, team)
        worksheet.cell(row_number, 2, player["player_name"])
        for column in range(1, len(headers) + 1):
            cell = worksheet.cell(row_number, column)
            _apply_font(cell, bold=column == 1, color=NAVY if column == 1 else INK)
            cell.alignment = Alignment(
                horizontal="left" if column <= 2 else "center",
                vertical="center",
                wrap_text=column >= 3,
            )
            cell.border = Border(bottom=Side(style="thin", color=LINE))
        worksheet.cell(row_number, 1).fill = _fill(TEAM_FILL)
        if team != previous_team:
            for column in range(1, len(headers) + 1):
                worksheet.cell(row_number, column).border = Border(
                    top=Side(style="medium", color=NAVY),
                    bottom=Side(style="thin", color=LINE),
                )
        previous_team = team

        for offset, week in enumerate(weeks, start=3):
            key = (team, str(player["player_id"]), int(week))
            game = cells.get(key)
            if game is None:
                continue
            cell = worksheet.cell(
                row_number,
                offset,
                f'{float(game["fantasy_points"]):.1f} pts\n'
                f'{float(game["defensive_snap_pct"]):.0%} snaps',
            )
            band = _rank_band(weekly_ranks[key], top_rank, second_rank)
            cell.fill = _fill(WEEKLY_FILLS[band])
            _apply_font(cell, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        player_cell = worksheet.cell(row_number, 2)
        season_band = _rank_band(season_ranks[_key(player)], top_rank, second_rank)
        if season_band != "below":
            player_cell.fill = _fill(SEASON_FILLS[season_band])
            _apply_font(player_cell, bold=True, color="FFFFFF" if season_band == "top" else "35264F")

        team_cell = worksheet.cell(row_number, 1)
        recent_band = _rank_band(recent_ranks[_key(player)], top_rank, second_rank)
        if recent_band != "below":
            team_cell.fill = _fill(RECENT_FILLS[recent_band])
            _apply_font(team_cell, bold=True, color="FFFFFF" if recent_band == "top" else "134E4A")
        else:
            team_cell.fill = _fill(TEAM_FILL)
            _apply_font(team_cell, bold=True, color=NAVY)

    last_row = len(players) + 3
    worksheet.auto_filter.ref = f"A3:{final_column}{last_row}"
    worksheet.row_dimensions[1].height = 24
    worksheet.row_dimensions[2].height = 18
    worksheet.row_dimensions[3].height = 22
    worksheet.column_dimensions["A"].width = 9
    worksheet.column_dimensions["B"].width = 24
    for column in range(3, len(headers) + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 11

    for cell in worksheet[1]:
        _apply_font(cell, bold=cell.column == 1, size=15 if cell.column == 1 else 10)
    _apply_font(worksheet["A2"], color=MUTED, size=9)
    worksheet["A2"].font = Font(name="Arial", size=9, italic=True, color=MUTED)
    for cell in worksheet[3]:
        cell.fill = _fill(NAVY)
        _apply_font(cell, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=Side(style="thin", color="FFFFFF"))
    for row in worksheet.iter_rows(min_row=4, max_row=last_row):
        worksheet.row_dimensions[row[0].row].height = 32


def _add_detail_sheet(workbook: Workbook, frame: pl.DataFrame) -> None:
    worksheet = workbook.create_sheet("Weekly detail")
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "C2"
    headers = [
        "Team",
        "Player",
        "Week",
        "Game ID",
        "Defensive snaps",
        "Defensive snap share",
        "Qualified at primary threshold",
        "Fantasy points",
        "Role state",
    ]
    columns = [
        "team",
        "player_name",
        "week",
        "game_id",
        "defensive_snaps",
        "defensive_snap_pct",
        "qualifies_full_time",
        "fantasy_points",
        "role_state",
    ]
    worksheet.append(headers)
    for row in frame.select(columns).sort(["team", "player_name", "week"]).iter_rows():
        worksheet.append(row)

    last_row = worksheet.max_row
    for cell in worksheet[1]:
        cell.fill = _fill(NAVY)
        _apply_font(cell, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            _apply_font(cell)
            cell.alignment = Alignment(vertical="center")
    for cell in worksheet["F"][1:]:
        cell.number_format = "0%"
    for cell in worksheet["E"][1:]:
        cell.number_format = "0"
    for cell in worksheet["H"][1:]:
        cell.number_format = "0.0"
    widths = {"A": 9, "B": 24, "C": 10, "D": 24, "E": 18, "F": 18, "G": 28, "H": 14, "I": 23}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    table = Table(displayName="WeeklyRoleDetail", ref=f"A1:I{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def write_workbook(
    season: int,
    threshold_frames: Sequence[tuple[float, pl.DataFrame]],
    destination: Path,
    *,
    top_rank: int = DEFAULT_TOP_RANK,
    second_rank: int = DEFAULT_SECOND_RANK,
    recent_week_count: int = DEFAULT_RECENT_WEEKS,
    regular_season_end: int = DEFAULT_REGULAR_SEASON_END,
) -> Path:
    if top_rank < 1 or second_rank <= top_rank:
        raise ValueError("Highlight ranks must satisfy 1 <= top rank < second rank")
    if recent_week_count < 1:
        raise ValueError("Recent-week count must be at least 1")
    if regular_season_end < 1:
        raise ValueError("Regular-season end must be at least 1")
    if not threshold_frames:
        raise ValueError("At least one threshold frame is required")

    workbook = Workbook()
    workbook.remove(workbook.active)
    for _, frame in threshold_frames:
        payload = _matrix_payload(
            frame,
            recent_week_count=recent_week_count,
            regular_season_end=regular_season_end,
        )
        _add_matrix_sheet(
            workbook,
            payload,
            season=season,
            top_rank=top_rank,
            second_rank=second_rank,
        )
    _add_detail_sheet(workbook, threshold_frames[0][1])
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    destination = destination.resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(f".{destination.stem}.tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(destination)
    return destination


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="python -m ndat.lb_workbook")
    parser.add_argument("--season", type=int, required=True)
    parser.add_argument("--refresh", action="store_true", help="refresh source data first")
    parser.add_argument("--primary-threshold", type=float, default=DEFAULT_PRIMARY_THRESHOLD)
    parser.add_argument("--comparison-threshold", type=float, default=DEFAULT_COMPARISON_THRESHOLD)
    parser.add_argument("--top-rank", type=int, default=DEFAULT_TOP_RANK)
    parser.add_argument("--second-rank", type=int, default=DEFAULT_SECOND_RANK)
    parser.add_argument("--recent-weeks", type=int, default=DEFAULT_RECENT_WEEKS)
    parser.add_argument("--regular-season-end", type=int, default=DEFAULT_REGULAR_SEASON_END)
    parser.add_argument("--output", type=Path)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    arguments = build_parser().parse_args(argv)
    config = DataConfig.from_project()
    manager = DataManager(config)
    refresh_results: list[dict[str, Any]] = []
    if arguments.refresh:
        for dataset in SOURCE_DATASETS:
            refresh_results.extend(manager.refresh(dataset, [arguments.season]))

    thresholds = [
        normalize_threshold(arguments.primary_threshold),
        normalize_threshold(arguments.comparison_threshold),
    ]
    if thresholds[0] == thresholds[1]:
        raise SystemExit("Primary and comparison thresholds must differ")

    threshold_frames: list[tuple[float, pl.DataFrame]] = []
    derived_paths: list[str] = []
    unsupported: set[str] = set()
    for threshold in thresholds:
        frame, path, missing_scoring = build_season(
            arguments.season, threshold=threshold, config=config
        )
        threshold_frames.append((threshold, frame))
        derived_paths.append(str(path))
        unsupported.update(missing_scoring)

    output = arguments.output or (
        config.derived_root
        / "lb_weekly_role"
        / f"season={arguments.season}"
        / f"linebacker_roles_{arguments.season}.xlsx"
    )
    workbook_path = write_workbook(
        arguments.season,
        threshold_frames,
        output,
        top_rank=arguments.top_rank,
        second_rank=arguments.second_rank,
        recent_week_count=arguments.recent_weeks,
        regular_season_end=arguments.regular_season_end,
    )
    print(
        json.dumps(
            {
                "season": arguments.season,
                "refreshed": refresh_results,
                "thresholds": thresholds,
                "derived_data": derived_paths,
                "workbook": str(workbook_path),
                "top_rank": arguments.top_rank,
                "second_rank": arguments.second_rank,
                "recent_weeks": arguments.recent_weeks,
                "regular_season_end": arguments.regular_season_end,
                "unsupported_scoring_categories": sorted(unsupported),
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
