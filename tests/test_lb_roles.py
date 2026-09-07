from __future__ import annotations

from pathlib import Path

import duckdb
import polars as pl
import pytest
from openpyxl import load_workbook

from ndat.config import DataConfig
from ndat.lb_roles import (
    IDP_SCORING,
    add_longitudinal_history,
    build_season,
    derive_weekly_roles,
    normalize_threshold,
    score_player_stats,
)
from ndat.lb_workbook import write_workbook


STAT_DEFAULTS = {
    field: 0
    for rule in IDP_SCORING.values()
    for field in rule.source_fields
}


def stats_row(**values: object) -> dict[str, object]:
    row: dict[str, object] = {
        "season": 2025,
        "week": 1,
        "game_id": "2025_01_BAL_BUF",
        "team": "BAL",
        "player_id": "00-0000001",
        "player_display_name": "Alex Example",
        **STAT_DEFAULTS,
    }
    row.update(values)
    return row


def source_frames(percentages: list[float] | None = None) -> tuple[pl.DataFrame, ...]:
    percentages = percentages or [0.849, 0.85, 0.851]
    count = len(percentages)
    snap = pl.DataFrame(
        {
            "season": [2025] * count,
            "week": list(range(1, count + 1)),
            "game_id": [f"2025_{week:02}_BAL_BUF" for week in range(1, count + 1)],
            "pfr_player_id": ["ExamAl00"] * count,
            "player": ["Alex Example"] * count,
            "team": ["BAL"] * count,
            "position": ["LB"] * count,
            "defense_snaps": [50.0] * count,
            "defense_pct": percentages,
        }
    )
    roster = pl.DataFrame(
        {
            "season": [2025],
            "week": [18],
            "pfr_id": ["ExamAl00"],
            "gsis_id": ["00-0000001"],
            "full_name": ["Alex Example"],
            "first_name": ["Alex"],
            "last_name": ["Example"],
            "position": ["LB"],
        }
    )
    stats = pl.DataFrame(
        [
            stats_row(
                week=week,
                game_id=f"2025_{week:02}_BAL_BUF",
                def_tackles_solo=week,
            )
            for week in range(1, count + 1)
        ]
    )
    return snap, roster, stats


def test_linebacker_selection_boundary_names_and_team_grouping() -> None:
    snap, roster, stats = source_frames()
    non_lb = snap.row(0, named=True) | {
        "pfr_player_id": "Other00",
        "player": "Other Player",
        "position": "DB",
    }
    result = derive_weekly_roles(pl.concat([snap, pl.DataFrame([non_lb])]), roster, stats)

    assert result.height == 3
    assert result["qualifies_full_time"].to_list() == [False, True, True]
    assert result["player_name"].unique().to_list() == ["Alex Example"]
    assert result["first_name"].unique().to_list() == ["Alex"]
    assert result["team"].unique().to_list() == ["BAL"]


def test_alternate_threshold_accepts_fraction_or_percentage() -> None:
    snap, roster, stats = source_frames([0.899, 0.9, 0.901])
    result = derive_weekly_roles(snap, roster, stats, threshold=90)

    assert normalize_threshold(90) == pytest.approx(0.9)
    assert result["qualifies_full_time"].to_list() == [False, True, True]
    with pytest.raises(ValueError):
        normalize_threshold(101)


@pytest.mark.parametrize(
    ("values", "expected"),
    [
        ({"def_sacks": 0.5}, 2),
        ({"def_tackles_solo": 3, "def_tackles_with_assist": 2}, 5),
        ({"def_punt_blocks": 1}, 5),
        ({"def_pat_blocks": 1}, 5),
        ({"def_fg_blocks": 1}, 5),
        ({"def_interceptions": 1}, 5),
        ({"fumble_recovery_opp": 1}, 2),
        ({"def_fumbles_forced": 1}, 4),
        ({"def_safeties": 1}, 8),
        ({"def_tackles_for_loss": 1}, 2),
        ({"def_sacks": 1, "def_tackles_for_loss": 1}, 6),
        ({"def_pass_defended": 1}, 1),
    ],
)
def test_each_supported_scoring_source(values: dict[str, object], expected: float) -> None:
    scored = score_player_stats(pl.DataFrame([stats_row(**values)]))
    assert scored["fantasy_points"][0] == pytest.approx(expected)


def test_total_tackles_do_not_double_count_assist_credit_and_nulls_are_zero() -> None:
    scored = score_player_stats(
        pl.DataFrame(
            [
                stats_row(
                    def_tackles_solo=4,
                    def_tackles_with_assist=3,
                    def_tackle_assists=99,  # Deliberately not part of the definition.
                    def_sacks=None,
                )
            ]
        )
    )
    assert scored["fantasy_points"][0] == 7


def test_longitudinal_acquisition_retention_loss_and_reacquisition() -> None:
    base = pl.DataFrame(
        {
            "season": [2025] * 6,
            "week": [1, 2, 3, 4, 5, 6],
            "game_id": [f"g{i}" for i in range(1, 7)],
            "player_id": ["p1"] * 6,
            "qualifies_full_time": [False, True, True, False, True, False],
        }
    )
    result = add_longitudinal_history(base)

    assert result["role_state"].to_list() == [
        "not_qualified",
        "acquired",
        "retained",
        "lost",
        "reacquired",
        "lost",
    ]
    assert result["first_qualifying_week"].unique().to_list() == [2]
    assert result["qualifying_weeks"][0].to_list() == [2, 3, 5]
    assert result["longest_qualifying_streak"].unique().to_list() == [2]
    assert result["reacquired"].to_list() == [False, False, False, False, True, False]


def test_stable_output_and_registered_duckdb_view(tmp_path: Path) -> None:
    config = DataConfig.from_project(project_root=tmp_path, data_root=tmp_path / "data")
    manager_paths = {
        name: config.source_root / name / "season=2025" / "data.parquet"
        for name in ("snap_counts", "rosters", "player_stats")
    }
    for path in manager_paths.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    for (name, path), frame in zip(manager_paths.items(), source_frames(), strict=True):
        frame.write_parquet(path)

    first, path, unsupported = build_season(2025, config=config)
    first_bytes = path.read_bytes()
    second, _, _ = build_season(2025, config=config)

    assert first.equals(second)
    assert path.read_bytes() == first_bytes
    assert unsupported == []
    with duckdb.connect(str(config.catalogue_path), read_only=True) as connection:
        assert connection.execute("SELECT count(*) FROM lb_weekly_role").fetchone() == (3,)


def test_workbook_navigation_values_and_highlights(tmp_path: Path) -> None:
    snap, roster, stats = source_frames()
    primary = derive_weekly_roles(snap, roster, stats, threshold=0.85)
    comparison = derive_weekly_roles(snap, roster, stats, threshold=0.80)
    destination = write_workbook(
        2025,
        [(0.85, primary), (0.80, comparison)],
        tmp_path / "linebacker_roles_2025.xlsx",
    )

    workbook = load_workbook(destination)
    assert workbook.sheetnames == ["85% Roles", "80% Roles", "Weekly detail"]
    assert workbook["85% Roles"].freeze_panes == "C4"
    assert workbook["80% Roles"].freeze_panes == "C4"
    assert workbook["Weekly detail"].freeze_panes == "C2"
    assert workbook["85% Roles"]["A1"].value == (
        "2025 linebacker snap-share roles at 85%"
    )
    assert workbook["85% Roles"]["C4"].value is None
    assert workbook["85% Roles"]["D4"].value == "2.0 pts\n85% snaps"
    assert workbook["85% Roles"]["B4"].fill.fgColor.rgb == "FF674EA7"
    assert workbook["85% Roles"]["A4"].fill.fgColor.rgb == "FF0F766E"
    assert workbook["85% Roles"].auto_filter.ref == "A3:E4"
    assert "WeeklyRoleDetail" in workbook["Weekly detail"].tables
